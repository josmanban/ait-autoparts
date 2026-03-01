from django.test import TestCase
from autoparts.services import AutoPartsImporter, CriticalErrorException, AutoPartsExporter
from autoparts.models import AutoPart, Category, Brand, Provider
from openpyxl import Workbook, load_workbook
from io import BytesIO


class AutoPartsImporterTestCase(TestCase):
    def setUp(self):
        self.category = Category.objects.get_or_create(name="Motor")
        self.brand = Brand.objects.get_or_create(name="MarcaA")
        self.provider = Provider.objects.get_or_create(name="ProvedorA")

    def test_import_valid_csv(self):
        with open('autoparts/tests/valid_autoparts.csv', 'rb') as file:
            importer = AutoPartsImporter(file)
            result = importer.import_from_csv()
            self.assertEqual(result['count'], 15)
            self.assertEqual(AutoPart.objects.count(), 15)
            self.assertEqual(len(result['success_records']), 15)
            self.assertEqual(len(result['failed_records']), 0)

    def test_import_csv_with_some_errors(self):
        with open('autoparts/tests/valid_with_some_errors.csv', 'rb') as file:
            importer = AutoPartsImporter(file)
            result = importer.import_from_csv()
            self.assertEqual(AutoPart.objects.count(), 2)
            self.assertTrue(AutoPart.objects.filter(code="ZZZ-001").exists())
            self.assertTrue(AutoPart.objects.filter(code="ZZZ-003").exists())
            self.assertEqual(result['count'], 2)
            self.assertIn("A valid number is required.", result['errors'][2]['unit_price'][0])
            self.assertEqual(len(result['success_records']), 2)
            self.assertEqual(len(result['failed_records']), 1)

    def test_import_csv_with_critical_errors(self):
        with open('autoparts/tests/fail_critical_errors.csv', 'rb') as file:
            importer = AutoPartsImporter(file)
            with self.assertRaises(CriticalErrorException) as context:
                importer.import_from_csv()
            self.assertIn("Import failed due to validation errors", str(context.exception))
            self.assertIn("Storage location must be in the format 'AA-12-03'", context.exception.summary["errors"][1]['storage_location'][0])
            self.assertIn("Code must be in the format 'AAA-001'", context.exception.summary["errors"][2]['code'][0])
            self.assertIn("A valid number is required.", context.exception.summary["errors"][2]['unit_price'][0])
            self.assertIn("Code must be in the format 'AAA-001'", context.exception.summary["errors"][3]['code'][0])
            self.assertEqual(len(context.exception.summary["success_records"]), 0)
            self.assertEqual(len(context.exception.summary["failed_records"]), 3)
            self.assertEqual(AutoPart.objects.count(), 0)

    def test_import_csv_with_critical_errors_global(self):
        with open('autoparts/tests/fail_critical_errors_global.csv', 'rb') as file:
            importer = AutoPartsImporter(file)
            with self.assertRaises(CriticalErrorException) as context:
                importer.import_from_csv()
            self.assertIn("Import failed due to validation errors", str(context.exception))
            self.assertIn("Duplicate codes found in the input data.", context.exception.summary["global_errors"])
            self.assertIn("Duplicate storage locations found in the input data.", context.exception.summary["global_errors"])
            self.assertIn("Duplicate name and brand combinations found in the input data.", context.exception.summary["global_errors"])
            self.assertEqual(len(context.exception.summary["success_records"]), 0)
            self.assertEqual(len(context.exception.summary["failed_records"]), 5)
            self.assertEqual(AutoPart.objects.count(), 0)

class AutoPartsExporterTestCase(TestCase):
    def setUp(self):
        self.category = Category.objects.get_or_create(name="Motor")[0]
        self.brand = Brand.objects.get_or_create(name="Toyota")[0]
        self.provider = Provider.objects.get_or_create(name="Provedor A")[0]

        AutoPart.objects.create(code="AAA-001", name="Part 1", description="Description 1", stock=10, min_stock=5, unit_price=100.00, category=self.category, brand=self.brand, provider=self.provider, storage_location="B-12-03")
        AutoPart.objects.create(code="AAA-002", name="Part 2", description="Description2", stock=20, min_stock=10, unit_price=200.00, category=self.category, brand=self.brand, provider=self.provider, storage_location="B-12-04")

    def test_export_to_stream(self):
        autoparts = AutoPart.objects.all()
        exporter = AutoPartsExporter(autoparts)
        stream = exporter.export_to_stream()
        workbook = load_workbook(BytesIO(stream))
        sheet = workbook.active
        self.assertEqual(sheet.title, "Autopartes")
        self.assertEqual(sheet.cell(row=1, column=1).value, "Codigo")
        self.assertEqual(sheet.cell(row=2, column=1).value, "AAA-001")
        self.assertEqual(sheet.cell(row=3, column=1).value, "AAA-002")
